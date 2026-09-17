#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
apply_correction_order_targets.py
'시정명령대상.xlsx' 파일의 대표 시정명령대상 정보를
facilities_cache.json, dispositions_cache.json 및 Supabase DB에 일괄 반영합니다.
"""
import openpyxl
import json
import requests
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.abspath(os.path.join(BASE_DIR, '..'))

EXCEL_FILE = os.path.join(PROJECT_DIR, '시정명령대상.xlsx')
FACILITIES_CACHE_FILE = os.path.join(PROJECT_DIR, 'facilities_cache.json')
DISPOSITIONS_CACHE_FILE = os.path.join(PROJECT_DIR, 'dispositions_cache.json')

SUPABASE_URL = 'https://vijiacxcmtfekbmegjlf.supabase.co'
SECRET_KEY = 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InZpamlhY3hjbXRmZWtibWVnamxmIiwicm9sZSI6InNlcnZpY2Vfcm9sZSIsImlhdCI6MTc4NTgyMzgyNiwiZXhwIjoyMTAxMzk5ODI2fQ.Noa3eCRZLGLp67fRYu4ZlsFC4_d2X1C7KxQ_g2_zP00'
HEADERS = {
    'apikey': SECRET_KEY,
    'Authorization': f'Bearer {SECRET_KEY}',
    'Content-Type': 'application/json',
    'Prefer': 'return=representation'
}

def main():
    print("=== [시정명령대상.xlsx DB 및 캐시 일괄 반영 시작] ===")
    
    # 1. 엑셀 데이터 로딩
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    sheet = wb.active
    
    mapping = {}
    for r in range(2, sheet.max_row + 1):
        seq = sheet.cell(r, 1).value
        name = sheet.cell(r, 2).value
        target = sheet.cell(r, 3).value
        key = sheet.cell(r, 4).value
        if key:
            k_str = str(key).strip()
            t_str = str(target).strip() if target else ''
            mapping[k_str] = {
                'seq': seq,
                'name': name,
                'target': t_str
            }
    print(f">> 엑셀 로딩 완료: 총 {len(mapping)}개 시설의 대표 시정명령대상 매핑")

    # 2. facilities_cache.json 업데이트
    with open(FACILITIES_CACHE_FILE, 'r', encoding='utf-8') as f:
        facilities = json.load(f)

    fac_updated_count = 0
    for fac in facilities:
        k = fac.get('facility_key')
        if k in mapping:
            fac['correction_order_target'] = mapping[k]['target']
            fac_updated_count += 1

    with open(FACILITIES_CACHE_FILE, 'w', encoding='utf-8') as f:
        json.dump(facilities, f, ensure_ascii=False, indent=2)
    print(f">> facilities_cache.json 업데이트 완료: {fac_updated_count}건 반영")

    # 3. dispositions_cache.json 업데이트
    with open(DISPOSITIONS_CACHE_FILE, 'r', encoding='utf-8') as f:
        dispositions = json.load(f)

    disp_updated_count = 0
    disp_ids_to_sync = []
    for d in dispositions:
        k = d.get('facility_key')
        if k in mapping and d.get('target_type') == '시설':
            d['correction_order'] = mapping[k]['target']
            disp_updated_count += 1
            if d.get('id'):
                disp_ids_to_sync.append((d['id'], mapping[k]['target']))

    with open(DISPOSITIONS_CACHE_FILE, 'w', encoding='utf-8') as f:
        json.dump(dispositions, f, ensure_ascii=False, indent=2)
    print(f">> dispositions_cache.json 업데이트 완료: {disp_updated_count}건 반영")

    # 4. Supabase DB 반영 (Direct PATCH)
    print(">> Supabase DB 동기화 진행 중...")
    db_fac_success = 0
    db_fac_fail = 0
    for k, info in mapping.items():
        try:
            r = requests.patch(
                f"{SUPABASE_URL}/rest/v1/facilities?facility_key=eq.{k}",
                headers=HEADERS,
                json={"correction_order_target": info['target']},
                timeout=5
            )
            if r.status_code in [200, 204]:
                db_fac_success += 1
            else:
                db_fac_fail += 1
        except Exception:
            db_fac_fail += 1

    print(f">> Supabase facilities PATCH 결과: 성공 {db_fac_success}건, 실패/미생성컬럼 {db_fac_fail}건")

    db_disp_success = 0
    for did, target in disp_ids_to_sync:
        try:
            r = requests.patch(
                f"{SUPABASE_URL}/rest/v1/dispositions?id=eq.{did}",
                headers=HEADERS,
                json={"correction_order": target},
                timeout=5
            )
            if r.status_code in [200, 204]:
                db_disp_success += 1
        except Exception:
            pass
    print(f">> Supabase dispositions(시설) PATCH 결과: 성공 {db_disp_success}건")

    print("=== [시정명령대상.xlsx 일괄 반영 완료] ===")

if __name__ == '__main__':
    main()
