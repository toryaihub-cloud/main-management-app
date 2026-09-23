# -*- coding: utf-8 -*-
"""
운영현황_1.xlsx의 추가 조사 122개 시설을 DB 및 캐시에 반영하는 스크립트
"""
import openpyxl
import datetime
import json
import sys
import os
import requests

# 프로젝트 루트 import 설정
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from crypto_utils import encrypt_data, decrypt_data

SUPABASE_URL = 'https://vijiacxcmtfekbmegjlf.supabase.co'
SECRET_KEY = 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InZpamlhY3hjbXRmZWtibWVnamxmIiwicm9sZSI6InNlcnZpY2Vfcm9sZSIsImlhdCI6MTc4NTgyMzgyNiwiZXhwIjoyMTAxMzk5ODI2fQ.Noa3eCRZLGLp67fRYu4ZlsFC4_d2X1C7KxQ_g2_zP00'
HEADERS = {
    'apikey': SECRET_KEY,
    'Authorization': 'Bearer ' + SECRET_KEY,
    'Content-Type': 'application/json',
    'Prefer': 'return=representation'
}

def to_int(val):
    if val is None:
        return 0
    s = str(val).strip()
    if not s or s == '-' or s.lower() == 'none':
        return 0
    try:
        return int(float(s))
    except Exception:
        return 0

def to_str(val):
    if val is None:
        return ''
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime('%Y-%m-%d')
    s = str(val).strip()
    if s.lower() == 'none':
        return ''
    return s

def main():
    print(">>> 1단계: 운영현황_1.xlsx 파싱 시작...")
    excel_path = os.path.join(BASE_DIR, '운영현황_1.xlsx')
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb.active
    headers = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]
    col_map = {h: idx + 1 for idx, h in enumerate(headers)}

    new_122_records = []
    sql_insert_values = []

    for r in range(2, sheet.max_row + 1):
        row_data = {h: sheet.cell(r, col_map[h]).value for h in headers}
        if not any(row_data.values()):
            continue

        key = to_str(row_data.get('KEY'))
        if not key:
            continue

        op_status = to_str(row_data.get('운영여부'))
        if op_status != '미운영':
            op_status = '정상운영'

        mgr_name = to_str(row_data.get('관리자'))
        mgr_contact = to_str(row_data.get('연락처'))

        inv_date = row_data.get('조사일')
        if isinstance(inv_date, (datetime.datetime, datetime.date)):
            inv_date_str = inv_date.strftime('%Y-%m-%d 00:00:00')
        else:
            inv_date_str = to_str(inv_date)

        enc_mgr = encrypt_data(mgr_name) if mgr_name else None
        enc_contact = encrypt_data(mgr_contact) if mgr_contact else None

        rec = {
            'facility_key': key,
            'facility_name': to_str(row_data.get('시설명')),
            'address_doro': to_str(row_data.get('주소[도로명]')),
            'investigator': to_str(row_data.get('조사자')),
            'investigation_date': inv_date_str,
            'parking_ground_cnt': to_int(row_data.get('설치면수(지상)')),
            'parking_underground_cnt': to_int(row_data.get('설치면수(지하)')),
            'parking_uninstalled_cnt': to_int(row_data.get('미설치면수')),
            'charger_installed_cnt': to_int(row_data.get('설치기수 합')),
            'charger_fast_cnt': to_int(row_data.get('설치기수(급속)')),
            'charger_slow_cnt': to_int(row_data.get('설치기수(완속)')),
            'charger_uninstalled_cnt': to_int(row_data.get('미설치기수')),
            'normal_operator_qty': to_str(row_data.get('정상운영 사업자(수량)')),
            'unoperated_cnt': to_int(row_data.get('미운영 기수')),
            'location': to_str(row_data.get('위치')),
            'unoperated_operator': to_str(row_data.get('미운영 사업자')),
            'initial_install_date': to_str(row_data.get('최초설치시기')),
            'operation_status': op_status,
            'unoperated_reason': to_str(row_data.get('미운영사유')),
            'unoperated_date': to_str(row_data.get('미운영시기')),
            'note': to_str(row_data.get('비고')),
            'complaint_and_plan': to_str(row_data.get('민원사항 및 향후계획')),
            'manager_name_encrypted': enc_mgr,
            'manager_contact_encrypted': enc_contact,
            'manager_name': mgr_name,
            'manager_contact': mgr_contact
        }
        new_122_records.append(rec)

        # SQL escape helper
        def sql_esc(v):
            if v is None:
                return "NULL"
            if isinstance(v, (int, float)):
                return str(v)
            return "'" + str(v).replace("'", "''") + "'"

        sql_insert_values.append(
            f"({sql_esc(rec['facility_key'])}, {sql_esc(rec['facility_name'])}, {sql_esc(rec['address_doro'])}, "
            f"{sql_esc(rec['investigator'])}, {sql_esc(rec['investigation_date'])}, {rec['parking_ground_cnt']}, "
            f"{rec['parking_underground_cnt']}, {rec['parking_uninstalled_cnt']}, {rec['charger_installed_cnt']}, "
            f"{rec['charger_fast_cnt']}, {rec['charger_slow_cnt']}, {rec['charger_uninstalled_cnt']}, "
            f"{sql_esc(rec['normal_operator_qty'])}, {rec['unoperated_cnt']}, {sql_esc(rec['location'])}, "
            f"{sql_esc(rec['unoperated_operator'])}, {sql_esc(rec['initial_install_date'])}, {sql_esc(rec['operation_status'])}, "
            f"{sql_esc(rec['unoperated_reason'])}, {sql_esc(rec['unoperated_date'])}, {sql_esc(rec['note'])}, "
            f"{sql_esc(rec['complaint_and_plan'])}, {sql_esc(rec['manager_name_encrypted'])}, {sql_esc(rec['manager_contact_encrypted'])})"
        )

    print(f"  [확인] 엑셀에서 {len(new_122_records)}개 시설 파싱 완료!")

    # >>> 2단계: 017_add_122_operations_data.sql 마이그레이션 파일 작성
    print(">>> 2단계: supabase/migrations/017_add_122_operations_data.sql 생성...")
    migration_file = os.path.join(BASE_DIR, 'supabase', 'migrations', '017_add_122_operations_data.sql')
    
    values_str = ",\n  ".join(sql_insert_values)
    sql_content = f"""-- 017_add_122_operations_data.sql
-- 운영현황_1.xlsx의 추가 조사 122개 시설 정보 등록

INSERT INTO public.operations (
    facility_key, facility_name, address_doro,
    investigator, investigation_date, parking_ground_cnt,
    parking_underground_cnt, parking_uninstalled_cnt, charger_installed_cnt,
    charger_fast_cnt, charger_slow_cnt, charger_uninstalled_cnt,
    normal_operator_qty, unoperated_cnt, location,
    unoperated_operator, initial_install_date, operation_status,
    unoperated_reason, unoperated_date, note,
    complaint_and_plan, manager_name_encrypted, manager_contact_encrypted
) VALUES
  {values_str};
"""
    with open(migration_file, 'w', encoding='utf-8') as f:
        f.write(sql_content)
    print(f"  [완료] {migration_file} 생성 완료!")

    # >>> 3단계: 기존 234건 캐시 로드 및 356건 병합
    print(">>> 3단계: 기존 operations_cache.json 로드 및 병합...")
    cache_path = os.path.join(BASE_DIR, 'operations_cache.json')
    old_cache = []
    if os.path.exists(cache_path):
        with open(cache_path, 'r', encoding='utf-8') as f:
            old_cache = json.load(f)

    # 기존 캐시 항목 복호화 필드(manager_name, manager_contact) 보충
    for item in old_cache:
        if 'manager_name' not in item or not item['manager_name']:
            item['manager_name'] = decrypt_data(item.get('manager_name_encrypted')) if item.get('manager_name_encrypted') else ''
        if 'manager_contact' not in item or not item['manager_contact']:
            item['manager_contact'] = decrypt_data(item.get('manager_contact_encrypted')) if item.get('manager_contact_encrypted') else ''

    old_keys = set(item['facility_key'] for item in old_cache if item.get('facility_key'))
    print(f"  기존 캐시 레코드 수: {len(old_cache)}건 (고유 키: {len(old_keys)})")

    # 병합
    merged_cache = list(old_cache)
    added_count = 0
    for rec in new_122_records:
        k = rec['facility_key']
        if k in old_keys:
            # 기존에 있으면 업데이트
            for idx, ex in enumerate(merged_cache):
                if ex.get('facility_key') == k:
                    merged_cache[idx] = rec
                    break
        else:
            merged_cache.append(rec)
            added_count += 1

    print(f"  병합 완료: 총 {len(merged_cache)}건 (신규 추가: {added_count}건)")

    # >>> 4단계: Supabase DB operations 테이블 적재
    print(">>> 4단계: Supabase DB에 적재/동기화...")
    # DB 기존 레코드 확인
    res = requests.get(f"{SUPABASE_URL}/rest/v1/operations?select=id,facility_key", headers=HEADERS)
    db_existing = res.json() if res.status_code == 200 else []
    print(f"  현재 Supabase DB 레코드 수: {len(db_existing)}건")

    db_key_map = {row['facility_key']: row['id'] for row in db_existing if row.get('facility_key')}

    # DB 전송용 데이터 (manager_name, manager_contact 평문 제외하고 DB 컬럼만)
    db_columns = [
        'facility_key', 'facility_name', 'address_doro', 'investigator', 'investigation_date',
        'parking_ground_cnt', 'parking_underground_cnt', 'parking_uninstalled_cnt',
        'charger_installed_cnt', 'charger_fast_cnt', 'charger_slow_cnt', 'charger_uninstalled_cnt',
        'normal_operator_qty', 'unoperated_cnt', 'location', 'unoperated_operator',
        'initial_install_date', 'operation_status', 'unoperated_reason', 'unoperated_date',
        'note', 'complaint_and_plan', 'manager_name_encrypted', 'manager_contact_encrypted'
    ]

    to_insert = []
    to_update = []

    for item in merged_cache:
        db_item = {col: item.get(col) for col in db_columns}
        k = item.get('facility_key')
        if k in db_key_map:
            to_update.append((db_key_map[k], db_item))
        else:
            to_insert.append(db_item)

    print(f"  DB INSERT 대상: {len(to_insert)}건, DB UPDATE 대상: {len(to_update)}건")

    # 배치 INSERT (50개씩)
    batch_size = 50
    for i in range(0, len(to_insert), batch_size):
        batch = to_insert[i:i + batch_size]
        res_ins = requests.post(f"{SUPABASE_URL}/rest/v1/operations", headers=HEADERS, json=batch, timeout=15)
        if res_ins.status_code in (200, 201):
            print(f"  [INSERT 성공] {i + 1} ~ {min(i + batch_size, len(to_insert))}행")
        else:
            print(f"  [INSERT 실패] status={res_ins.status_code}, error={res_ins.text}")
            return False

    # UPDATE 대상 처리
    for did, db_item in to_update:
        res_upd = requests.patch(f"{SUPABASE_URL}/rest/v1/operations?id=eq.{did}", headers=HEADERS, json=db_item, timeout=10)
        if res_upd.status_code not in (200, 204):
            print(f"  [UPDATE 실패] id={did}, error={res_upd.text}")

    # 최종 DB 건수 확인
    res_final = requests.get(f"{SUPABASE_URL}/rest/v1/operations?select=id", headers=HEADERS)
    final_db_count = len(res_final.json()) if res_final.status_code == 200 else 0
    print(f"  [확인] Supabase DB operations 최종 적재 건수: {final_db_count}건")

    # >>> 5단계: operations_cache.json 업데이트
    print(">>> 5단계: operations_cache.json 갱신...")
    with open(cache_path, 'w', encoding='utf-8') as f:
        json.dump(merged_cache, f, ensure_ascii=False, indent=2)
    print(f"  [완료] operations_cache.json에 {len(merged_cache)}건 저장 완료!")

    print("\n=================================================================")
    print(f"  [성공] 운영현황 122개 추가 및 총 {len(merged_cache)}개 시설 동기화 완료!")
    print("=================================================================")
    return True

if __name__ == '__main__':
    ok = main()
    sys.exit(0 if ok else 1)
